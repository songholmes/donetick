#!/usr/bin/env python3
"""Import Leona's timetable into Donetick as grouped recurring chores.

Dry-run by default. Use --apply with DONETICK_USERNAME and DONETICK_PASSWORD
to create missing chores through the Donetick HTTP API.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


DEFAULT_EXCEL = r"C:\Users\Songholmes\Downloads\Leona's Time Table.xlsx"
DEFAULT_BASE_URL = "http://localhost:2021"
PROJECT_NAME = "Leona Household Timetable"
PROJECT_DESCRIPTION = "Imported grouped schedule from Leona's Time Table.xlsx"
SHEET_NAME = "Version 2"
TIMEZONE = "Asia/Singapore"

LABELS: dict[str, str] = {
    "Weekday": "#2563EB",
    "Weekend": "#7C3AED",
    "Morning": "#F59E0B",
    "Baby": "#EC4899",
    "Cleaning": "#10B981",
    "Cooking": "#EF4444",
    "Pets": "#84CC16",
    "Errands": "#06B6D4",
    "Evening": "#6366F1",
    "Optional": "#94A3B8",
}

WEEKDAY_DAYS = ["monday", "tuesday", "wednesday", "thursday", "friday"]
WEEKEND_DAYS = ["saturday", "sunday"]


@dataclass(frozen=True)
class ChoreSpec:
    key: str
    name: str
    days: list[str]
    due_time: str
    source_range: str
    value_cell: str
    labels: list[str]


CHORE_SPECS: list[ChoreSpec] = [
    ChoreSpec("leona-v2-weekday-0600", "Morning opening routine", WEEKDAY_DAYS, "06:00", "Version 2!B3:F3", "B3", ["Weekday", "Morning", "Pets", "Cooking", "Baby", "Optional"]),
    ChoreSpec("leona-v2-weekday-0700", "Baby morning care", WEEKDAY_DAYS, "07:00", "Version 2!B4:F4", "B4", ["Weekday", "Morning", "Baby"]),
    ChoreSpec("leona-v2-weekday-0730", "School drop-off", WEEKDAY_DAYS, "07:30", "Version 2!B5:F5", "B5", ["Weekday", "Baby"]),
    ChoreSpec("leona-v2-weekday-0800", "House reset, cleaning, and laundry", WEEKDAY_DAYS, "08:00", "Version 2!B6:F6", "B6", ["Weekday", "Cleaning", "Optional"]),
    ChoreSpec("leona-v2-weekday-1115", "Prepare lunch", WEEKDAY_DAYS, "11:15", "Version 2!B7:F7", "B7", ["Weekday", "Cooking"]),
    ChoreSpec("leona-v2-weekday-1200", "Lunch cleanup and rest block", WEEKDAY_DAYS, "12:00", "Version 2!B8:F8", "B8", ["Weekday", "Cooking", "Cleaning"]),
    ChoreSpec("leona-v2-weekday-1500", "Dinner prep and cooking", WEEKDAY_DAYS, "15:00", "Version 2!B11:F11", "B11", ["Weekday", "Cooking", "Errands", "Cleaning"]),
    ChoreSpec("leona-v2-weekday-1700", "Fetch baby", WEEKDAY_DAYS, "17:00", "Version 2!B12:F12", "B12", ["Weekday", "Baby"]),
    ChoreSpec("leona-v2-weekday-1730", "Baby dinner / play support", WEEKDAY_DAYS, "17:30", "Version 2!B13:F13", "B13", ["Weekday", "Baby"]),
    ChoreSpec("leona-v2-weekday-1845", "Evening closing routine", WEEKDAY_DAYS, "18:45", "Version 2!B14:F14", "B14", ["Weekday", "Evening", "Baby", "Pets", "Cleaning"]),
    ChoreSpec("leona-v2-monday-1330", "Bedsheet, bed, sofa, vacuum cleanup", ["monday"], "13:30", "Version 2!B9", "B9", ["Weekday", "Cleaning"]),
    ChoreSpec("leona-v2-tuesday-1330", "Bathrooms, mats, slippers, shoes", ["tuesday"], "13:30", "Version 2!C9", "C9", ["Weekday", "Cleaning"]),
    ChoreSpec("leona-v2-wednesday-1330", "Shower glass, mirrors, windows, doors, shoe cabinet", ["wednesday"], "13:30", "Version 2!D9", "D9", ["Weekday", "Cleaning"]),
    ChoreSpec("leona-v2-thursday-1330", "Microwave, baby equipment, high chair, stroller", ["thursday"], "13:30", "Version 2!E9", "E9", ["Weekday", "Cleaning", "Cooking", "Baby"]),
    ChoreSpec("leona-v2-friday-1330", "Bins, fridge, grocery list", ["friday"], "13:30", "Version 2!F9", "F9", ["Weekday", "Cleaning", "Errands"]),
    ChoreSpec("leona-v2-saturday-0600", "Weekend morning opening routine", ["saturday"], "06:00", "Version 2!G3", "G3", ["Weekend", "Morning", "Pets", "Cooking", "Optional"]),
    ChoreSpec("leona-v2-saturday-0700", "Wet market groceries", ["saturday"], "07:00", "Version 2!G4:G5", "G4", ["Weekend", "Morning", "Errands"]),
    ChoreSpec("leona-v2-sunday-0630", "Sunday morning opening routine", ["sunday"], "06:30", "Version 2!H3:H4", "H3", ["Weekend", "Morning", "Pets", "Cooking", "Optional"]),
    ChoreSpec("leona-v2-sunday-0730", "Sunday baby morning care", ["sunday"], "07:30", "Version 2!H5", "H5", ["Weekend", "Morning", "Baby"]),
    ChoreSpec("leona-v2-weekend-0800", "Weekend plan-dependent house or baby support", WEEKEND_DAYS, "08:00", "Version 2!G6:H8", "G6", ["Weekend", "Cleaning", "Baby", "Optional"]),
    ChoreSpec("leona-v2-weekend-1330", "Baby afternoon tea if home", WEEKEND_DAYS, "13:30", "Version 2!G9:H10", "G9", ["Weekend", "Baby", "Optional"]),
    ChoreSpec("leona-v2-weekend-1500", "Weekend dinner prep / normal work", WEEKEND_DAYS, "15:00", "Version 2!G11:H11", "G11", ["Weekend", "Cooking", "Cleaning"]),
    ChoreSpec("leona-v2-weekend-1700", "Walk Miya if early", WEEKEND_DAYS, "17:00", "Version 2!G12:H12", "G12", ["Weekend", "Pets", "Optional"]),
    ChoreSpec("leona-v2-weekend-1730", "Weekend evening baby / normal work block", WEEKEND_DAYS, "17:30", "Version 2!G13:H14", "G13", ["Weekend", "Evening", "Baby", "Pets", "Cleaning"]),
]


def merged_value(ws: Worksheet, coordinate: str) -> str:
    cell = ws[coordinate]
    if cell.value is not None:
        return str(cell.value).strip()
    for cell_range in ws.merged_cells.ranges:
        if coordinate in cell_range:
            value = ws.cell(cell_range.min_row, cell_range.min_col).value
            return "" if value is None else str(value).strip()
    return ""


def split_subtasks(text: str) -> list[str]:
    subtasks: list[str] = []
    for raw_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("-"):
            item = line.lstrip("-").strip()
            if item:
                subtasks.append(item)
        elif subtasks:
            subtasks[-1] = f"{subtasks[-1]} {line}".strip()
        else:
            subtasks.append(line)
    return subtasks


def parse_hhmm(value: str) -> time:
    return datetime.strptime(value, "%H:%M").time()


def next_due_iso(days: list[str], hhmm: str, tz_name: str, include_today: bool = True) -> str:
    loc = ZoneInfo(tz_name)
    now = datetime.now(loc)
    wanted = set(days)
    target_time = parse_hhmm(hhmm)
    for offset in range(0, 8):
        candidate_date = now.date() + timedelta(days=offset)
        candidate = datetime.combine(candidate_date, target_time, loc)
        if candidate.strftime("%A").lower() in wanted and (candidate > now or (include_today and offset == 0)):
            return candidate.isoformat(timespec="seconds")
    raise RuntimeError(f"Unable to calculate next due date for days={days}, time={hhmm}")


def load_preview(excel_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(excel_path, data_only=True)
    if workbook.sheetnames != [SHEET_NAME]:
        raise RuntimeError(f"Expected exactly one sheet named {SHEET_NAME!r}; found {workbook.sheetnames!r}")
    ws = workbook[SHEET_NAME]
    preview: list[dict[str, Any]] = []
    for spec in CHORE_SPECS:
        raw_text = merged_value(ws, spec.value_cell)
        subtasks = split_subtasks(raw_text)
        if not subtasks:
            raise RuntimeError(f"No subtasks extracted for {spec.key} from {spec.value_cell}")
        description = f"Imported from Leona's Time Table.xlsx | {spec.source_range} | importKey={spec.key}"
        preview.append(
            {
                "importKey": spec.key,
                "name": spec.name,
                "days": spec.days,
                "dueTime": spec.due_time,
                "nextDueDate": next_due_iso(spec.days, spec.due_time, TIMEZONE),
                "sourceRange": spec.source_range,
                "description": description,
                "labels": spec.labels,
                "subtasks": subtasks,
            }
        )
    return preview


def print_preview(preview: list[dict[str, Any]]) -> None:
    print(f"Workbook: {DEFAULT_EXCEL}")
    print(f"Sheet: {SHEET_NAME}")
    print(f"Chores planned: {len(preview)}")
    print()
    for idx, item in enumerate(preview, start=1):
        day_text = ",".join(day[:3].title() for day in item["days"])
        labels = ", ".join(item["labels"])
        print(f"{idx:02d}. {item['name']} [{day_text} {item['dueTime']}]")
        print(f"    key: {item['importKey']}")
        print(f"    labels: {labels}")
        print(f"    subtasks: {len(item['subtasks'])}")
        for subtask in item["subtasks"][:5]:
            print(f"      - {subtask}")
        if len(item["subtasks"]) > 5:
            print(f"      ... +{len(item['subtasks']) - 5} more")
        print()


def json_request(method: str, base_url: str, path: str, token: str | None = None, payload: Any | None = None) -> Any:
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    data = json.dumps(payload).encode("utf-8") if payload is not None else None
    request = urllib.request.Request(base_url.rstrip("/") + path, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            body = response.read().decode("utf-8")
            return json.loads(body) if body else None
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"{method} {path} failed with HTTP {exc.code}: {body}") from exc


def unwrap_items(response: Any) -> list[dict[str, Any]]:
    if isinstance(response, list):
        return response
    if isinstance(response, dict):
        for key in ("res", "data", "items"):
            value = response.get(key)
            if isinstance(value, list):
                return value
    raise RuntimeError(f"Could not unwrap list response: {response!r}")


def login(base_url: str, username: str, password: str) -> str:
    response = json_request("POST", base_url, "/api/v1/auth/login", payload={"username": username, "password": password})
    token = response.get("access_token") or response.get("token")
    if not token:
        raise RuntimeError(f"Login response did not include a token: {response!r}")
    return str(token)


def get_profile(base_url: str, token: str) -> dict[str, Any]:
    response = json_request("GET", base_url, "/api/v1/users/profile", token=token)
    return response.get("res", response) if isinstance(response, dict) else response


def ensure_project(base_url: str, token: str) -> int:
    projects = unwrap_items(json_request("GET", base_url, "/api/v1/projects", token=token))
    for project in projects:
        if project.get("name") == PROJECT_NAME:
            return int(project["id"])
    created = json_request(
        "POST",
        base_url,
        "/api/v1/projects",
        token=token,
        payload={"name": PROJECT_NAME, "description": PROJECT_DESCRIPTION, "color": "#8B5CF6", "icon": "calendar"},
    )
    project = created.get("res") if isinstance(created, dict) else None
    if not project or "id" not in project:
        raise RuntimeError(f"Project creation did not return an ID: {created!r}")
    return int(project["id"])


def ensure_labels(base_url: str, token: str) -> dict[str, int]:
    existing = unwrap_items(json_request("GET", base_url, "/api/v1/labels", token=token))
    by_name = {str(label.get("name")): int(label["id"]) for label in existing if label.get("id") is not None}
    for name, color in LABELS.items():
        if name in by_name:
            continue
        created = json_request("POST", base_url, "/api/v1/labels", token=token, payload={"name": name, "color": color})
        label = created.get("res") if isinstance(created, dict) else None
        if not label or "id" not in label:
            raise RuntimeError(f"Label creation for {name!r} did not return an ID: {created!r}")
        by_name[name] = int(label["id"])
    return by_name


def existing_import_keys(base_url: str, token: str) -> set[str]:
    chores = unwrap_items(json_request("GET", base_url, "/api/v1/chores/?includeArchived=true", token=token))
    keys: set[str] = set()
    for chore in chores:
        match = re.search(r"importKey=([^\s|]+)", chore.get("description") or "")
        if match:
            keys.add(match.group(1))
    return keys


def create_chore(base_url: str, token: str, item: dict[str, Any], user_id: int, project_id: int, label_ids: dict[str, int]) -> int:
    payload = {
        "name": item["name"],
        "frequencyType": "days_of_the_week",
        "frequency": 1,
        "frequencyMetadata": {
            "days": item["days"],
            "time": item["nextDueDate"],
            "timezone": TIMEZONE,
            "weekPattern": "every_week",
        },
        "nextDueDate": item["nextDueDate"],
        "isRolling": False,
        "assignedTo": user_id,
        "assignees": [{"userId": user_id}],
        "assignStrategy": "keep_last_assigned",
        "notification": False,
        "notificationMetadata": None,
        "labelsV2": [{"id": label_ids[name]} for name in item["labels"]],
        "priority": 0,
        "description": item["description"],
        "subTasks": [{"id": -(idx + 1), "name": name, "orderId": idx + 1} for idx, name in enumerate(item["subtasks"])],
        "requireApproval": False,
        "isPrivate": False,
        "projectId": project_id,
    }
    created = json_request("POST", base_url, "/api/v1/chores/", token=token, payload=payload)
    if isinstance(created, dict) and "res" in created:
        return int(created["res"])
    raise RuntimeError(f"Chore creation for {item['name']!r} did not return an ID: {created!r}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Import Leona timetable as grouped Donetick chores.")
    parser.add_argument("--excel", default=DEFAULT_EXCEL)
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--apply", action="store_true", help="Actually create missing project/labels/chores in Donetick.")
    parser.add_argument("--username", default=os.environ.get("DONETICK_USERNAME"))
    parser.add_argument("--password", default=os.environ.get("DONETICK_PASSWORD"))
    parser.add_argument("--preview-json", default="")
    args = parser.parse_args()

    preview = load_preview(Path(args.excel))
    if len(preview) != 24:
        raise RuntimeError(f"Expected 24 grouped chores; generated {len(preview)}")
    if args.preview_json:
        Path(args.preview_json).write_text(json.dumps(preview, indent=2, ensure_ascii=False), encoding="utf-8")
    print_preview(preview)

    if not args.apply:
        print("Dry run only. Re-run with --apply to create missing chores.")
        return 0
    if not args.username or not args.password:
        raise RuntimeError("DONETICK_USERNAME and DONETICK_PASSWORD are required for --apply.")

    token = login(args.base_url, args.username, args.password)
    profile = get_profile(args.base_url, token)
    user_id = int(profile["id"])
    project_id = ensure_project(args.base_url, token)
    label_ids = ensure_labels(args.base_url, token)
    seen_keys = existing_import_keys(args.base_url, token)

    created: list[tuple[str, int]] = []
    skipped: list[str] = []
    for item in preview:
        if item["importKey"] in seen_keys:
            skipped.append(item["name"])
            continue
        chore_id = create_chore(args.base_url, token, item, user_id, project_id, label_ids)
        created.append((item["name"], chore_id))

    print()
    print(f"Created chores: {len(created)}")
    for name, chore_id in created:
        print(f"  + #{chore_id}: {name}")
    print(f"Skipped existing chores: {len(skipped)}")
    for name in skipped:
        print(f"  = {name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
