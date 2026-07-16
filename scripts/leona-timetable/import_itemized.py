#!/usr/bin/env python3
"""Import Leona's timetable into Donetick as one chore per bullet item.

Dry-run by default. Use --apply with DONETICK_USERNAME and DONETICK_PASSWORD
to create missing chores through the Donetick HTTP API.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Literal
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from import_grouped import (  # noqa: E402
    CHORE_SPECS,
    DEFAULT_BASE_URL,
    DEFAULT_EXCEL,
    LABELS,
    SHEET_NAME,
    TIMEZONE,
    existing_imported_chores,
    get_profile,
    json_request,
    login,
    merged_value,
    next_due_iso,
    parse_hhmm,
    split_subtasks,
    unwrap_items,
    set_due_date,
)


PROJECT_NAME = "Leona Household Timetable - Itemized"
PROJECT_DESCRIPTION = "Imported itemized schedule from Leona's Time Table.xlsx"
IMPORT_PREFIX = "leona-v2-itemized"
EXPECTED_ITEMIZED_COUNT = 71

Cadence = Literal["block", "interval"]


@dataclass(frozen=True)
class ItemizedChore:
    import_key: str
    name: str
    days: list[str]
    start_time: str
    due_time: str
    next_due_date: str
    source_window: str
    source_range: str
    source_bullet: str
    labels: list[str]
    frequency_type: str
    frequency: int
    frequency_metadata: dict[str, Any]


SPECIAL_SPLITS: dict[str, list[dict[str, Any]]] = {
    "Clean the master bedroom: Reset the bed suit every day, clean the toilet every other day": [
        {"name": "Reset the bed suit", "cadence": "block"},
        {"name": "Clean the toilet", "cadence": "interval", "frequency": 2, "unit": "days"},
    ],
}


def slugify(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return normalized[:60].strip("-") or "item"


def unique_labels(labels: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for label in labels:
        if label not in seen:
            seen.add(label)
            result.append(label)
    return result


def labels_for_item(block_labels: list[str], name: str) -> list[str]:
    labels = [label for label in block_labels if label != "Optional"]
    lower = name.lower()
    if any(marker in lower for marker in ("optional", "option ", "if needed", "as needed", "if baby", "if we ", "unless ", "if you finish")):
        labels.append("Optional")
    if any(marker in lower for marker in ("clean", "wash", "vacuum", "vaccum", "mop", "tidy", "fold", "trash", "bedsheet", "bed suit", "toilet")):
        labels.append("Cleaning")
    if any(marker in lower for marker in ("breakfast", "lunch", "dinner", "cook", "kitchen", "dish", "dishes", "bowl", "groceries", "ingredients")):
        labels.append("Cooking")
    if any(marker in lower for marker in ("baby", "uniform", "school", "bottle", "sterilize", "stroller", "high chair")):
        labels.append("Baby")
    if any(marker in lower for marker in ("miya", "rabbit", "pet")):
        labels.append("Pets")
    if any(marker in lower for marker in ("buy", "market", "ntuc", "grocery list", "groceries")):
        labels.append("Errands")
    if any(marker in lower for marker in ("sleep", "next day", "back to your room")):
        labels.append("Evening")
    return unique_labels(labels)


def split_source_range(source_range: str) -> str:
    if "!" not in source_range:
        return source_range
    return source_range.split("!", 1)[1]


def parse_time_range(value: str) -> tuple[str, str]:
    match = re.search(r"(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})", value)
    if not match:
        raise RuntimeError(f"Could not parse timetable range from {value!r}")
    start, end = match.groups()
    return normalize_hhmm(start), normalize_hhmm(end)


def normalize_hhmm(value: str) -> str:
    hour, minute = value.split(":", 1)
    return f"{int(hour):02d}:{minute}"


def due_time_from_source_range(ws: Any, source_range: str, display_start_time: str | None = None) -> tuple[str, str]:
    """Return the end time for the source range and the displayed source window.

    Donetick has one due datetime, so itemized chores use the end of the source
    timetable block as the overdue threshold. For merged timetable blocks, the
    source range can span multiple rows; the final row determines the end time.
    """

    cell_range = split_source_range(source_range)
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    del min_col, max_col

    first_range = merged_value(ws, f"A{min_row}")
    last_range = merged_value(ws, f"A{max_row}")
    if not first_range or not last_range:
        raise RuntimeError(f"Could not read time window for {source_range}")

    start_time_from_sheet, _ = parse_time_range(first_range)
    _, end_time = parse_time_range(last_range)
    display_start = normalize_hhmm(display_start_time) if display_start_time else start_time_from_sheet
    return end_time, f"{display_start} - {end_time}"


def next_interval_due_iso(days: list[str], hhmm: str, tz_name: str) -> str:
    """Use the next due date from the source block as the first interval due date."""
    return next_due_iso(days, hhmm, tz_name)


def interval_frequency_metadata(next_due_date: str, unit: str) -> dict[str, Any]:
    return {
        "time": next_due_date,
        "timezone": TIMEZONE,
        "unit": unit,
    }


def block_frequency_metadata(days: list[str], next_due_date: str) -> dict[str, Any]:
    return {
        "days": days,
        "time": next_due_date,
        "timezone": TIMEZONE,
        "weekPattern": "every_week",
    }


def expand_bullet(block_key: str, bullet_index: int, bullet: str, block_labels: list[str]) -> list[dict[str, Any]]:
    if bullet in SPECIAL_SPLITS:
        expanded = []
        for split_index, split in enumerate(SPECIAL_SPLITS[bullet], start=1):
            expanded.append(
                {
                    "name": split["name"],
                    "source_bullet": bullet,
                    "key_suffix": f"{bullet_index:02d}-{split_index:02d}-{slugify(split['name'])}",
                    "cadence": split["cadence"],
                    "frequency": split.get("frequency", 1),
                    "unit": split.get("unit"),
                    "labels": labels_for_item(block_labels, split["name"]),
                }
            )
        return expanded
    return [
        {
            "name": bullet,
            "source_bullet": bullet,
            "key_suffix": f"{bullet_index:02d}-{slugify(bullet)}",
            "cadence": "block",
            "frequency": 1,
            "unit": None,
            "labels": labels_for_item(block_labels, bullet),
        }
    ]


def build_itemized_preview(excel_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(excel_path, data_only=True)
    if workbook.sheetnames != [SHEET_NAME]:
        raise RuntimeError(f"Expected exactly one sheet named {SHEET_NAME!r}; found {workbook.sheetnames!r}")
    ws = workbook[SHEET_NAME]

    preview: list[dict[str, Any]] = []
    for spec in CHORE_SPECS:
        raw_text = merged_value(ws, spec.value_cell)
        bullets = split_subtasks(raw_text)
        if not bullets:
            raise RuntimeError(f"No bullet items extracted for {spec.key} from {spec.value_cell}")
        due_time, source_window = due_time_from_source_range(ws, spec.source_range, spec.due_time)

        block_slug = spec.key.removeprefix("leona-v2-")
        for bullet_index, bullet in enumerate(bullets, start=1):
            expanded_items = expand_bullet(block_slug, bullet_index, bullet, spec.labels)
            for expanded in expanded_items:
                import_key = f"{IMPORT_PREFIX}-{block_slug}-{expanded['key_suffix']}"
                if expanded["cadence"] == "interval":
                    next_due = next_interval_due_iso(spec.days, due_time, TIMEZONE)
                    frequency_type = "interval"
                    frequency = int(expanded["frequency"])
                    frequency_metadata = interval_frequency_metadata(next_due, str(expanded["unit"]))
                else:
                    next_due = next_due_iso(spec.days, due_time, TIMEZONE)
                    frequency_type = "days_of_the_week"
                    frequency = 1
                    frequency_metadata = block_frequency_metadata(spec.days, next_due)
                description = (
                    f"Imported itemized from Leona's Time Table.xlsx | {spec.source_range} | "
                    f"sourceWindow={source_window} | sourceBullet={expanded['source_bullet']} | "
                    f"importKey={import_key}"
                )
                preview.append(
                    {
                        "importKey": import_key,
                        "name": expanded["name"],
                        "days": spec.days,
                        "startTime": spec.due_time,
                        "dueTime": due_time,
                        "nextDueDate": next_due,
                        "sourceWindow": source_window,
                        "sourceRange": spec.source_range,
                        "sourceBullet": expanded["source_bullet"],
                        "labels": expanded["labels"],
                        "frequencyType": frequency_type,
                        "frequency": frequency,
                        "frequencyMetadata": frequency_metadata,
                        "description": description,
                    }
                )
    return preview


def print_preview(preview: list[dict[str, Any]]) -> None:
    print(f"Workbook: {DEFAULT_EXCEL}")
    print(f"Sheet: {SHEET_NAME}")
    print(f"Project: {PROJECT_NAME}")
    print(f"Chores planned: {len(preview)}")
    print()
    for idx, item in enumerate(preview, start=1):
        day_text = ",".join(day[:3].title() for day in item["days"])
        labels = ", ".join(item["labels"])
        cadence = item["frequencyType"]
        if cadence == "interval":
            cadence = f"interval/{item['frequency']} {item['frequencyMetadata'].get('unit')}"
        print(f"{idx:02d}. {item['name']} [{day_text} {item['startTime']}-{item['dueTime']} due@{item['dueTime']} {cadence}]")
        print(f"    key: {item['importKey']}")
        print(f"    source: {item['sourceRange']}")
        print(f"    window: {item['sourceWindow']}")
        print(f"    labels: {labels}")
        print()


def ensure_project(base_url: str, token: str) -> int:
    projects = unwrap_items(json_request("GET", base_url, "/api/v1/projects", token=token))
    for project in projects:
        if project.get("name") == PROJECT_NAME:
            return int(project["id"])
    created = json_request(
        "POST",
        base_url,
        "/api/v1/projects",
        token=token,
        payload={"name": PROJECT_NAME, "description": PROJECT_DESCRIPTION, "color": "#0EA5E9", "icon": "checklist"},
    )
    project = created.get("res") if isinstance(created, dict) else None
    if not project or "id" not in project:
        raise RuntimeError(f"Project creation did not return an ID: {created!r}")
    return int(project["id"])


def ensure_labels(base_url: str, token: str) -> dict[str, int]:
    existing = unwrap_items(json_request("GET", base_url, "/api/v1/labels", token=token))
    by_name = {str(label.get("name")): int(label["id"]) for label in existing if label.get("id") is not None}
    for name, color in LABELS.items():
        if name in by_name:
            continue
        created = json_request("POST", base_url, "/api/v1/labels", token=token, payload={"name": name, "color": color})
        label = created.get("res") if isinstance(created, dict) else None
        if not label or "id" not in label:
            raise RuntimeError(f"Label creation for {name!r} did not return an ID: {created!r}")
        by_name[name] = int(label["id"])
    return by_name


def existing_itemized_chores(base_url: str, token: str, project_id: int) -> dict[str, dict[str, Any]]:
    imported = existing_imported_chores(base_url, token)
    return {
        key: chore
        for key, chore in imported.items()
        if key.startswith(f"{IMPORT_PREFIX}-") and int(chore.get("projectId") or 0) == project_id
    }


def delete_chore(base_url: str, token: str, chore_id: int) -> None:
    json_request("DELETE", base_url, f"/api/v1/chores/{chore_id}", token=token)


def delete_chore_history(base_url: str, token: str, chore_id: int, history_id: int) -> None:
    json_request("DELETE", base_url, f"/api/v1/chores/{chore_id}/history/{history_id}", token=token)


def delete_reschedule_histories(base_url: str, token: str, chore_ids: list[int]) -> list[tuple[int, int]]:
    deleted: list[tuple[int, int]] = []
    for chore_id in chore_ids:
        histories = unwrap_items(json_request("GET", base_url, f"/api/v1/chores/{chore_id}/history", token=token))
        for history in histories:
            if int(history.get("status") or -1) != 6:
                continue
            history_id = int(history["id"])
            delete_chore_history(base_url, token, chore_id, history_id)
            deleted.append((chore_id, history_id))
    return deleted


def reset_existing_itemized(base_url: str, token: str, project_id: int) -> list[tuple[str, int, str]]:
    existing = existing_itemized_chores(base_url, token, project_id)
    deleted: list[tuple[str, int, str]] = []
    for key, chore in sorted(existing.items(), key=lambda entry: int(entry[1]["id"])):
        chore_id = int(chore["id"])
        name = str(chore.get("name") or "")
        delete_chore(base_url, token, chore_id)
        deleted.append((key, chore_id, name))
    return deleted


def create_chore(base_url: str, token: str, item: dict[str, Any], user_id: int, project_id: int, label_ids: dict[str, int]) -> int:
    payload = {
        "name": item["name"],
        "frequencyType": item["frequencyType"],
        "frequency": item["frequency"],
        "frequencyMetadata": item["frequencyMetadata"],
        "nextDueDate": item["nextDueDate"],
        "isRolling": False,
        "assignedTo": user_id,
        "assignees": [{"userId": user_id}],
        "assignStrategy": "keep_last_assigned",
        "notification": False,
        "notificationMetadata": None,
        "labelsV2": [{"id": label_ids[name]} for name in item["labels"]],
        "priority": 0,
        "description": item["description"],
        "subTasks": [],
        "requireApproval": False,
        "isPrivate": False,
        "projectId": project_id,
    }
    created = json_request("POST", base_url, "/api/v1/chores/", token=token, payload=payload)
    if isinstance(created, dict) and "res" in created:
        chore_id = int(created["res"])
        set_due_date(base_url, token, chore_id, item["nextDueDate"])
        return chore_id
    raise RuntimeError(f"Chore creation for {item['name']!r} did not return an ID: {created!r}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Import Leona timetable as itemized Donetick chores.")
    parser.add_argument("--excel", default=DEFAULT_EXCEL)
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--apply", action="store_true", help="Actually create missing project/labels/chores in Donetick.")
    parser.add_argument(
        "--reset-existing",
        action="store_true",
        help="With --apply, delete existing Leona itemized chores in the target project before recreating them.",
    )
    parser.add_argument("--username", default=os.environ.get("DONETICK_USERNAME"))
    parser.add_argument("--password", default=os.environ.get("DONETICK_PASSWORD"))
    parser.add_argument("--preview-json", default="")
    args = parser.parse_args()

    preview = build_itemized_preview(Path(args.excel))
    if len(preview) != EXPECTED_ITEMIZED_COUNT:
        raise RuntimeError(f"Expected {EXPECTED_ITEMIZED_COUNT} itemized chores; generated {len(preview)}")
    if args.preview_json:
        Path(args.preview_json).write_text(json.dumps(preview, indent=2, ensure_ascii=False), encoding="utf-8")
    print_preview(preview)

    if args.reset_existing and not args.apply:
        raise RuntimeError("--reset-existing requires --apply.")
    if not args.apply:
        print("Dry run only. Re-run with --apply to create missing itemized chores.")
        return 0
    if not args.username or not args.password:
        raise RuntimeError("DONETICK_USERNAME and DONETICK_PASSWORD are required for --apply.")

    token = login(args.base_url, args.username, args.password)
    profile = get_profile(args.base_url, token)
    user_id = int(profile["id"])
    project_id = ensure_project(args.base_url, token)
    label_ids = ensure_labels(args.base_url, token)
    deleted: list[tuple[str, int, str]] = []
    if args.reset_existing:
        deleted = reset_existing_itemized(args.base_url, token, project_id)
    seen_chores = existing_itemized_chores(args.base_url, token, project_id)

    created: list[tuple[str, int]] = []
    skipped: list[str] = []
    for item in preview:
        existing = seen_chores.get(item["importKey"])
        if existing:
            if not existing.get("nextDueDate"):
                set_due_date(args.base_url, token, int(existing["id"]), item["nextDueDate"])
            skipped.append(item["name"])
            continue
        chore_id = create_chore(args.base_url, token, item, user_id, project_id, label_ids)
        created.append((item["name"], chore_id))

    cleaned_histories: list[tuple[int, int]] = []
    if args.reset_existing and created:
        cleaned_histories = delete_reschedule_histories(args.base_url, token, [chore_id for _, chore_id in created])

    print()
    if args.reset_existing:
        print(f"Deleted existing itemized chores: {len(deleted)}")
        for _, chore_id, name in deleted:
            print(f"  - #{chore_id}: {name}")
        print(f"Deleted import-created reschedule histories: {len(cleaned_histories)}")
    print(f"Created chores: {len(created)}")
    for name, chore_id in created:
        print(f"  + #{chore_id}: {name}")
    print(f"Skipped existing chores: {len(skipped)}")
    for name in skipped:
        print(f"  = {name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
