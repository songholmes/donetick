from __future__ import annotations

import unittest

from openpyxl import Workbook

from import_itemized import due_time_from_source_range, parse_time_range


class EndTimeParsingTests(unittest.TestCase):
    def test_single_row_uses_row_end_time(self) -> None:
        workbook = Workbook()
        ws = workbook.active
        ws["A6"] = "8:00 - 11:15"

        due_time, source_window = due_time_from_source_range(ws, "Version 2!B6:F6")

        self.assertEqual(due_time, "11:15")
        self.assertEqual(source_window, "08:00 - 11:15")

    def test_multi_row_source_uses_final_row_end_time(self) -> None:
        workbook = Workbook()
        ws = workbook.active
        ws["A6"] = "8:00 - 11:15"
        ws["A7"] = "11:15 - 11:45"
        ws["A8"] = "12:00 - 13:30"

        due_time, source_window = due_time_from_source_range(ws, "Version 2!G6:H8")

        self.assertEqual(due_time, "13:30")
        self.assertEqual(source_window, "08:00 - 13:30")

    def test_merged_final_time_cell_is_resolved(self) -> None:
        workbook = Workbook()
        ws = workbook.active
        ws["A9"] = "13:30 - 15:00"
        ws.merge_cells("A9:A10")

        due_time, source_window = due_time_from_source_range(ws, "Version 2!G9:H10")

        self.assertEqual(due_time, "15:00")
        self.assertEqual(source_window, "13:30 - 15:00")

    def test_time_range_normalizes_single_digit_hour(self) -> None:
        self.assertEqual(parse_time_range("6:00 - 7:00"), ("06:00", "07:00"))


if __name__ == "__main__":
    unittest.main()
